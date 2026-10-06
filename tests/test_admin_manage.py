"""
Админ-панель (записка заказчика 02.10.2026): сотрудник вручную, импорт продуктов и страховых случаев из Excel.
Всё — на ВРЕМЕННОЙ копии базы (tests/tmpdb.py), имена и телефоны выдуманные.

Запуск из корня проекта:
    set PYTHONIOENCODING=utf-8
    sandbox\\.venv\\Scripts\\python.exe tests\\test_admin_manage.py
"""
import asyncio
import io
import json
import os
import sys
from datetime import date, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))
os.environ.pop("SURVEYOR_DEV", None)                  # вход по сессии, как на Railway

import openpyxl                                        # noqa: E402

from tmpdb import temp_db                              # noqa: E402
from app import auth, calibration, db                  # noqa: E402
from app.main import app                               # noqa: E402

PASSED = []
TOK = {}


def ok(name, cond, hint=""):
    assert cond, f"ПРОВАЛ: {name} {hint}"
    PASSED.append(name)
    print("  ✓", name)


def call(method, path, body=None, who=None, raw=None, ctype="application/json", token=None):
    payload = raw if raw is not None else (json.dumps(body, ensure_ascii=False).encode() if body is not None else b"")
    path, _, query = path.partition("?")
    hdrs = [(b"host", b"test"), (b"content-type", ctype.encode()), (b"content-length", str(len(payload)).encode())]
    tok = token or (TOK[who] if who else None)
    if tok:
        hdrs.append((b"authorization", f"Bearer {tok}".encode()))
    scope = {"type": "http", "asgi": {"version": "3.0"}, "http_version": "1.1", "method": method,
             "scheme": "http", "path": path, "raw_path": path.encode(), "root_path": "",
             "query_string": query.encode(), "headers": hdrs,
             "client": ("203.0.113.9", 0), "server": ("test", 80)}
    out = {"status": None, "chunks": [], "headers": []}

    async def receive():
        return {"type": "http.request", "body": payload, "more_body": False}

    async def send(msg):
        if msg["type"] == "http.response.start":
            out["status"] = msg["status"]
            out["headers"] = msg.get("headers") or []
        elif msg["type"] == "http.response.body":
            out["chunks"].append(msg.get("body") or b"")

    asyncio.run(app(scope, receive, send))
    data = b"".join(out["chunks"])
    try:
        return out["status"], json.loads(data.decode("utf-8")), out["headers"]
    except (ValueError, UnicodeDecodeError):
        return out["status"], data, out["headers"]


def upload(path, data, who, name="файл.xlsx"):
    b = "----admin-manage-boundary"
    body = (f"--{b}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"{name}\"\r\n"
            "Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n").encode() \
        + data + f"\r\n--{b}--\r\n".encode()
    return call("POST", path, raw=body, who=who, ctype=f"multipart/form-data; boundary={b}")


def xlsx(rows) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def sid_of(headers) -> str:
    for k, v in headers:
        if k == b"set-cookie" and v.startswith(b"sid="):
            return v.split(b";")[0][4:].decode()
    return ""


def setup():
    with db.tx() as con:
        for login, role in (("тест-админ-панель", "админ"), ("тест-сотрудник-панель", "сотрудник")):
            cur = con.execute(
                "INSERT INTO users (login, full_name, phone, role, branch, password_hash, salt, status, created_at)"
                " VALUES (?,?,?,?,?,?,?,?,?)",
                (login, "Тест " + role, "", role, "тест", "x" * 64, "0" * 32, auth.STATUS_ACTIVE, db.now()))
            u = db.rows(con, "SELECT * FROM users WHERE id=?", cur.lastrowid)[0]
            TOK[role], _ = auth.create_session(con, u, ip="203.0.113.9", user_agent="test")


# --------------------------------------------------------------------------- #

def check_staff():
    print("Сотрудник вручную")
    body = {"full_name": "Тестов Тест Тестович", "position": "менеджер", "department": "Тестовый департамент",
            "unit": "Отдел проверок", "branch": "Тестовый филиал", "phone": "+998 90 000-00-91",
            "login": "Test.Manual1", "temp_password": "Vremenny-1"}
    st, out, _ = call("POST", "/tg/users/manual", body, who="сотрудник")
    ok("не-админ не заводит сотрудника → 403", st == 403, out)
    st, out, _ = call("POST", "/tg/users/manual", body, who="админ")
    ok("админ заводит сотрудника → 200", st == 200, out)
    u = out["user"]
    ok("телефон нормализован, логин в нижнем регистре",
       u["phone"] == "+998900000091" and u["login"] == "test.manual1", u)
    ok("активен, роль сотрудник, отдел сохранён, нужна смена пароля",
       u["status"] == auth.STATUS_ACTIVE and u["role"] == "сотрудник" and u["unit"] == "Отдел проверок"
       and u["must_change_password"] is True, u)
    ok("пароль задан админом — в ответе не повторяется", "temp_password" not in out)
    uid = u["id"]

    st, out, _ = call("POST", "/tg/users/manual", body | {"phone": "+998900000092"}, who="админ")
    ok("тот же логин → 409", st == 409, out)
    st, out, _ = call("POST", "/tg/users/manual", body | {"login": "test.manual2"}, who="админ")
    ok("тот же телефон → 409", st == 409, out)
    st, out, _ = call("POST", "/tg/users/manual", body | {"login": "test.manual3", "phone": "12345"}, who="админ")
    ok("неверный телефон → 422", st == 422, out)
    st, out, _ = call("POST", "/tg/users/manual", body | {"login": "test.manual3", "phone": "+998900000093",
                                                         "temp_password": "123"}, who="админ")
    ok("короткий временный пароль → 422", st == 422, out)
    st, out, _ = call("POST", "/tg/users/manual", body | {"login": "test.manual4", "phone": "+998900000094",
                                                         "temp_password": ""}, who="админ")
    ok("без пароля — сервер создаёт временный и показывает один раз",
       st == 200 and len(out.get("temp_password") or "") >= 8, out)
    other_uid = out["user"]["id"]

    with db.tx() as con:
        h = db.rows(con, "SELECT password_hash, salt FROM users WHERE id=?", uid)[0]
    ok("пароль хранится хэшем PBKDF2", auth.check_password("Vremenny-1", h["password_hash"], h["salt"])
       and "Vremenny" not in h["password_hash"])

    # вход с временным паролем
    st, out, hd = call("POST", "/auth/login", {"login": "test.manual1", "password": "Vremenny-1"})
    tok = sid_of(hd)
    ok("вход с временным паролем: ответ требует смены", st == 200 and out.get("must_change_password") is True
       and out.get("next") == "PUT /auth/password" and tok, out)
    st, out, _ = call("GET", "/tg/users", token=tok)
    ok("до смены пароля остальное закрыто → 403 must_change_password",
       st == 403 and out.get("must_change_password") is True, out)
    st, out, _ = call("GET", "/auth/me", token=tok)
    ok("/auth/me доступен и показывает флаг", st == 200 and out["must_change_password"] is True, out)
    st, out, _ = call("GET", "/tg/me", token=tok)
    ok("/tg/me с временным паролем: флаг и пустое меню", st == 200 and out.get("must_change_password") is True
       and out["nav"] == [] and out["rights"] == [], out)
    st, out, _ = call("POST", "/auth/admin-request", {}, token=tok)
    ok("публичный /auth/admin-request с временным паролем → 403 must_change_password",
       st == 403 and out.get("must_change_password") is True, out)
    st, out, _ = call("PUT", "/auth/password", {"old_password": "неверный", "new_password": "Novyi-parol-1"},
                      token=tok)
    ok("смена с неверным текущим → 403", st == 403, out)
    st, out, _ = call("PUT", "/auth/password", {"old_password": "Vremenny-1", "new_password": "Vremenny-1"},
                      token=tok)
    ok("новый пароль равен временному → 422", st == 422, out)
    st, out, _ = call("PUT", "/auth/password", {"old_password": "Vremenny-1", "new_password": "Novyi-parol-1"},
                      token=tok)
    ok("смена пароля → 200", st == 200 and out["must_change_password"] is False, out)
    st, out, _ = call("GET", "/tg/users", token=tok)
    ok("после смены пароля сессия работает", st == 200, out)
    st, _, _ = call("POST", "/auth/login", {"login": "test.manual1", "password": "Vremenny-1"})
    ok("временный пароль больше не подходит → 401", st == 401)
    st, out, _ = call("POST", "/auth/login", {"login": "test.manual1", "password": "Novyi-parol-1"})
    ok("вход с новым паролем без требования смены", st == 200 and not out.get("must_change_password"), out)

    # правка карточки
    st, out, _ = call("PUT", f"/tg/users/{uid}", {"position": "директор", "unit": "Новый отдел",
                                                  "phone": "998900000095"}, who="сотрудник")
    ok("не-админ не правит карточку → 403", st == 403, out)
    st, out, _ = call("PUT", f"/tg/users/{uid}", {"position": "директор", "unit": "Новый отдел",
                                                  "phone": "998900000095"}, who="админ")
    ok("правка карточки → поля изменены", st == 200 and set(out["changed"]) == {"position", "unit", "phone"}
       and out["user"]["phone"] == "+998900000095", out)
    st, out, _ = call("PUT", f"/tg/users/{uid}", {"phone": "+998900000094"}, who="админ")
    ok("правка: чужой телефон → 409", st == 409, out)
    st, out, _ = call("PUT", "/tg/users/999999", {"unit": "x"}, who="админ")
    ok("правка несуществующего → 404", st == 404, out)

    # сброс пароля
    st, out, _ = call("POST", f"/tg/users/{uid}/reset-password", who="сотрудник")
    ok("не-админ не сбрасывает пароль → 403", st == 403, out)
    st, out, _ = call("POST", f"/tg/users/{uid}/reset-password", who="админ")
    new_pw = out.get("temp_password") or ""
    ok("сброс: новый временный пароль один раз", st == 200 and len(new_pw) >= 8 and out["must_change_password"])
    st, _, _ = call("GET", "/auth/me", token=tok)
    ok("сброс закрыл прежние сессии → 401", st == 401)
    st, out, _ = call("POST", "/auth/login", {"login": "test.manual1", "password": new_pw})
    ok("вход после сброса снова требует смены", st == 200 and out.get("must_change_password") is True, out)

    with db.tx() as con:
        logs = db.rows(con, "SELECT detail FROM audit WHERE entity IN (?,?)", f"user:{uid}", f"user:{other_uid}")
    text = " ".join(r["detail"] or "" for r in logs)
    ok("журнал: телефон только маской, без ФИО и паролей",
       logs and "900000091" not in text and "900000095" not in text and "+998*******" in text
       and "Тестов" not in text and "Vremenny" not in text and new_pw not in text, text[:300])
    st, out, _ = call("GET", "/tg/users", who="админ")
    me = next(x for x in out["items"] if x["id"] == uid)
    ok("список пользователей отдаёт отдел", me["unit"] == "Новый отдел", me)


def check_templates():
    print("Шаблоны")
    for path in ("/reference/products/template.xlsx", "/claims/template.xlsx"):
        st, data, hd = call("GET", path, who="админ")
        ok(f"{path} → xlsx с PK", st == 200 and isinstance(data, bytes) and data[:2] == b"PK")
        wb = openpyxl.load_workbook(io.BytesIO(data))
        ok(f"{path}: есть лист «Инструкция»", "Инструкция" in wb.sheetnames)
    st, data, _ = call("GET", "/claims/template.xlsx", who="админ")
    head = [c.value for c in openpyxl.load_workbook(io.BytesIO(data)).worksheets[0][1]]
    ok("шаблон убытков без колонок ФИО", not any("фио" in str(h).lower() or "страхователь" in str(h).lower()
                                                 for h in head), head)


def check_products() -> str:
    print("Импорт продуктов")
    today = date.today()
    tday = today.strftime("%d.%m.%Y")
    later = (today + timedelta(days=30)).strftime("%d.%m.%Y")
    from app.act_engine import part_rates_from_text
    with db.tx() as con:
        prods = db.rows(con, "SELECT p.code, p.rate_text, (SELECT COUNT(*) FROM product_classes c"
                             " WHERE c.product_code = p.code) n FROM products p"
                             " WHERE p.pricing_mode='ставка' AND p.code NOT LIKE '99%' ORDER BY p.code")
        cands = [r["code"] for r in prods if r["n"] == 1 and r["code"] != "0807"
                 and set(part_rates_from_text(r["rate_text"], [])) == {"*"}]
        exist, simple = cands[0], cands[1]
        multi = next(r["code"] for r in prods if r["n"] > 1)
        r0807 = db.rows(con, "SELECT rate_text FROM products WHERE code='0807'")[0]["rate_text"]
        versions0 = con.execute("SELECT COUNT(*) FROM tariff_versions").fetchone()[0]
        classes0 = con.execute("SELECT COUNT(*) FROM product_classes").fetchone()[0]
    head = ["Код", "Название", "Классы (через запятую)", "Базовая ставка %", "Минимальная ставка %",
            "Тип ставки (годовая|фиксированная)", "Дата вступления", "Статус (тест|действует)"]
    rows = [head,
            ["9901", "Тестовый продукт импорта", "8", 0.5, 0.3, "годовая", tday, "тест"],
            ["9902", "Ошибочный продукт", "8,99x", "abc", None, "годовая", None, "тест"],
            [exist, None, None, None, 0.2, "фиксированная", later, None],
            ["9903", "Продукт задним числом", "8", 0.4, None, "годовая", "01.01.2020", "тест"],
            ["9904", "Многоклассовый с базовой", "8,9", 0.5, None, "годовая", tday, "тест"],
            ["0807", None, None, 0.07, None, "годовая", later, None],
            [multi, None, None, 0.3, None, "годовая", tday, None],
            [simple, None, None, 0.77, None, "годовая", tday, None]]
    data = xlsx(rows)
    st, out, _ = upload("/reference/products/import", data, "сотрудник")
    ok("не-админ не импортирует продукты → 403", st == 403, out)
    st, out, _ = upload("/reference/products/import", b"not a zip at all", "админ")
    ok("файл без сигнатуры PK → 400", st == 400, out)
    big = openpyxl.Workbook()
    big.active.append(head)
    big.active.cell(row=5000, column=1, value="x")
    buf = io.BytesIO()
    big.save(buf)
    st, out, _ = upload("/reference/products/import", buf.getvalue(), "админ")
    ok("лист больше предела строк (пустые тоже считаются) → 422", st == 422 and "Строк" in out["detail"], out)
    wide = openpyxl.Workbook()
    wide.active.append(head + ["лишняя"] * 70)
    buf = io.BytesIO()
    wide.save(buf)
    st, out, _ = upload("/reference/products/import", buf.getvalue(), "админ")
    ok("больше 64 колонок → 422", st == 422 and "Колонок" in out["detail"], out)

    st, out, _ = upload("/reference/products/import", data, "админ", "продукты.xlsx")
    ok("предпросмотр → 200 и token", st == 200 and out.get("token"), out)
    it = {x["product_code"]: x for x in out["items"]}
    ok("ok_count 3, error_count 5", out["ok_count"] == 3 and out["error_count"] == 5,
       [(x["product_code"], x["errors"]) for x in out["items"]])
    ok("новый одноклассовый продукт с базовой сегодня → create", it["9901"]["action"] == "create"
       and it["9901"]["classes"] == ["8"] and it["9901"]["rate_pct"] == 0.5, it["9901"])
    ok("ошибочная строка: класс и ставка", it["9902"]["action"] == "skip"
       and any("99x" in e for e in it["9902"]["errors"]) and any("Базовая" in e for e in it["9902"]["errors"]),
       it["9902"])
    ok("существующий продукт: минимальная с будущей датой → update",
       it[exist]["action"] == "update" and it[exist]["rate_type"] == "fixed", it[exist])
    ok("дата в прошлом → ошибка", any("Дата" in e for e in it["9903"]["errors"]), it["9903"])
    ok("базовая для многоклассового нового продукта → ошибка",
       any("нескольких классов" in e for e in it["9904"]["errors"]), it["9904"])
    ok("базовая с будущей датой (0807) → ошибка", any("будущей датой" in e for e in it["0807"]["errors"]),
       it["0807"])
    ok("базовая для существующего многоклассового → ошибка",
       any("нескольких классов" in e for e in it[multi]["errors"]), it[multi])
    ok("базовая для простого одноклассового сегодня → update", it[simple]["action"] == "update", it[simple])

    st, out2, _ = call("POST", "/reference/products/import/apply", {"token": out["token"]}, who="сотрудник")
    ok("не-админ не применяет → 403", st == 403, out2)
    st, res, _ = call("POST", "/reference/products/import/apply", {"token": out["token"]}, who="админ")
    ok("apply → создан 1, обновлено 2", st == 200 and res["created"] == 1 and res["updated"] == 2, res)
    st, again, _ = call("POST", "/reference/products/import/apply", {"token": out["token"]}, who="админ")
    ok("токен одноразовый → 404", st == 404, again)
    from app import act, min_rates
    with db.tx() as con:
        p = db.rows(con, "SELECT * FROM products WHERE code='9901'")[0]
        ps = db.rows(con, "SELECT rate_text FROM products WHERE code=?", simple)[0]
        r0807_after = db.rows(con, "SELECT rate_text FROM products WHERE code='0807'")[0]["rate_text"]
        pcs = [r["class_code"] for r in db.rows(con, "SELECT class_code FROM product_classes WHERE product_code="
                                                     "'9901' ORDER BY part_no")]
        versions1 = con.execute("SELECT COUNT(*) FROM tariff_versions").fetchone()[0]
        prv = db.rows(con, "SELECT p.*, v.effective_from FROM product_rate_versions p JOIN tariff_versions v"
                           " ON v.id = p.tariff_version_id WHERE p.product_code IN ('9901', ?)", simple)
        mr = db.rows(con, "SELECT * FROM min_rate_versions WHERE product_code IN ('9901', ?)", exist)
        classes1 = con.execute("SELECT COUNT(*) FROM product_classes").fetchone()[0]
        bad = db.rows(con, "SELECT code FROM products WHERE code IN ('9902','9903','9904')")
        audit = db.rows(con, "SELECT 1 FROM audit WHERE action='продукт' AND entity='9901'")
        fut = min_rates.on_date(con, exist, today + timedelta(days=31))
        own = {r["tariff_version_id"] for r in db.rows(con, "SELECT tariff_version_id FROM product_rate_versions"
                                                            " UNION SELECT tariff_version_id FROM min_rate_versions")}
        tv_act = act.tariff_version(con, "компания")
        tv_calc = min_rates.current_version(con, "компания")
        notes = [r["note"] for r in mr]
    ok("продукт создан в статусе «тест», текст тарифа «0,5%»", p["status"] == "тест" and p["rate_text"] == "0,5%",
       p)
    ok("простой одноклассовый: текст тарифа стал «0,77%»", ps["rate_text"] == "0,77%", ps)
    ok("0807: текст тарифа не изменён импортом", r0807_after == r0807 == "0,05%", (r0807, r0807_after))
    ok("классы продукта 8", pcs == ["8"], pcs)
    ok("apply создал версии тарифов (2 базовые + 2 минимума)", versions1 - versions0 == 4, (versions0, versions1))
    ok("базовая ставка — версией с сегодняшней датой", len(prv) == 2
       and all(r["effective_from"] == today.isoformat() for r in prv), prv)
    ok("минимальные ставки — через min_rates (источник import, примечание пустое)",
       len(mr) == 2 and all(r["source"] == "import" for r in mr) and notes == ["", ""], mr)
    ok("будущая минимальная ставка существующего продукта: 0,2% fixed с даты",
       fut and fut["pct"] == 0.2 and fut["rate_type"] == "fixed", fut)
    ok("ошибочные строки не записаны", not bad, bad)
    ok("из справочника ничего не удалено", classes1 == classes0 + 1, (classes0, classes1))
    ok("журнал «продукт» — защита от перезаписи refsync", bool(audit))
    ok("акт и /calculate берут действующую версию уровня, а не версию правки продукта",
       tv_act and tv_act == tv_calc and tv_act not in own, (tv_act, tv_calc))

    st, out, _ = upload("/reference/products/import", data, "админ")
    it = {x["product_code"]: x for x in out["items"]}
    ok("повторная загрузка: продукт без изменений → skip", it["9901"]["action"] == "skip", it["9901"])
    return "9901"


def check_claims(code: str):
    print("Импорт страховых случаев")
    t = date.today()
    d = lambda n: (t - timedelta(days=n)).strftime("%d.%m.%Y")   # noqa: E731
    head = ["Номер дела", "Код продукта", "Класс", "Номер договора", "Филиал", "Регион", "Вид объекта",
            "Дата события", "Дата заявления", "Дата выплаты", "Причина", "Заявлено (сум)", "Выплачено (сум)",
            "Статус (заявлен|урегулирован|отказ|в работе)", "Страховая сумма", "Премия", "ФИО страхователя",
            "ИНН страхователя"]
    rows = [head,
            ["ТЕСТ-У-1", code, "8", "ТД-1", "Тестовый филиал", "Тестовый регион", "склад", d(100), d(98), d(60),
             "пожар", 20_000_000, 15_000_000, "урегулирован", 500_000_000, 3_000_000, "Выдуманов Выдуман",
             "123456789"],
            ["ТЕСТ-У-1", code, "8", "ТД-1", "", "", "склад", d(90), d(88), None, "пожар", 1000, 0, "заявлен",
             None, None, "Выдуманов Выдуман"],
            ["ТЕСТ-У-2", code, "", "ТД-2", "", "", "", d(50), d(49), None, "кража", 5_000_000, None, "в работе",
             100_000_000, 1_000_000, "Придумова Придума"],
            ["ТЕСТ-У-3", code, "8", "ТД-3", "", "", "склад", d(40), None, None, "залив", 1000, None, "закрыт",
             None, None, ""],
            ["ТЕСТ-У-4", code, "8", "ТД-4", "", "", "склад", d(30), None, None, "залив", 1000, -5, "отказ",
             None, None, ""],
            ["ТЕСТ-У-5", "0000", "8", "ТД-5", "", "", "склад", (t + timedelta(days=5)).strftime("%d.%m.%Y"),
             None, None, "", 1000, None, "заявлен", None, None, ""]]
    data = xlsx(rows)
    st, out, _ = upload("/claims/import", data, "сотрудник")
    ok("не-админ не загружает убытки → 403", st == 403, out)
    st, out, _ = upload("/claims/import", data, "админ", "убытки.xlsx")
    ok("предпросмотр убытков → 200 и token", st == 200 and out.get("token"), out)
    ok("ok 2, ошибок 4 (дубль, статус, минус, продукт/дата)", out["ok_count"] == 2 and out["error_count"] == 4,
       [(x["row"], x["errors"]) for x in out["items"]])
    by_row = {x["row"]: x for x in out["items"]}
    ok("дубль номера дела в файле — ошибка", any("повтор" in e for e in by_row[3]["errors"]), by_row[3])
    ok("статус не из списка — ошибка", any("Статус" in e for e in by_row[5]["errors"]), by_row[5])
    ok("сумма меньше нуля — ошибка", any("меньше нуля" in e for e in by_row[6]["errors"]), by_row[6])
    ok("нет продукта и дата в будущем — ошибки", any("Код продукта" in e for e in by_row[7]["errors"])
       and any("будущем" in e for e in by_row[7]["errors"]), by_row[7])
    ok("класс не указан — взят первый класс продукта", by_row[4]["class_code"] == "8"
       and by_row[4]["warnings"], by_row[4])
    ok("колонки ФИО и ИНН пропущены с предупреждением",
       out["pd_columns_ignored"] == ["ФИО страхователя", "ИНН страхователя"]
       and any("персональными" in w for w in out["warnings"]), out["warnings"])
    ok("ФИО нет в предпросмотре", "Выдуманов" not in json.dumps(out, ensure_ascii=False))

    st, res, _ = call("POST", "/claims/import/apply", {"token": out["token"]}, who="админ")
    ok("apply → создано 2, пропущено 4", st == 200 and res["created"] == 2 and res["skipped"] == 4, res)
    with db.tx() as con:
        b = db.rows(con, "SELECT * FROM claim_batches WHERE id=?", res["batch_id"])[0]
        c1 = db.rows(con, "SELECT * FROM claims WHERE external_no='ТЕСТ-У-1'")
        dump = json.dumps(db.rows(con, "SELECT * FROM claims WHERE batch_id=?", res["batch_id"]), ensure_ascii=False)
    ok("батч: файл, строки, кто, когда", b["file_name"] == "убытки.xlsx" and b["rows"] == 6
       and b["loaded_by"] == "тест-админ-панель" and b["loaded_at"], b)
    ok("новые колонки claims заполнены", len(c1) == 1 and c1[0]["product_code"] == code
       and c1[0]["contract_no"] == "ТД-1" and c1[0]["branch"] == "Тестовый филиал"
       and c1[0]["sum_insured"] == 500_000_000 and c1[0]["premium"] == 3_000_000, c1)
    ok("ФИО и ИНН в базу не попали", "Выдуманов" not in dump and "Придумова" not in dump
       and "123456789" not in dump)

    # повторная загрузка того же номера — обновление
    rows2 = [head[:16], ["ТЕСТ-У-1", code, "8", "ТД-1", "Тестовый филиал", "Тестовый регион", "склад", d(100),
                         d(98), d(55), "пожар", 20_000_000, 18_000_000, "урегулирован", 500_000_000, 3_000_000]]
    st, out, _ = upload("/claims/import", xlsx(rows2), "админ", "убытки-2.xlsx")
    ok("повтор номера дела в базе → update", st == 200 and out["items"][0]["action"] == "update", out)
    st, res, _ = call("POST", "/claims/import/apply", {"token": out["token"]}, who="админ")
    with db.tx() as con:
        c1 = db.rows(con, "SELECT * FROM claims WHERE external_no='ТЕСТ-У-1'")
    ok("обновлено, не дубль", st == 200 and res["updated"] == 1 and len(c1) == 1 and c1[0]["paid"] == 18_000_000,
       (res, c1))

    st, s, _ = call("GET", "/claims/summary", who="сотрудник")
    ok("сводка убытков не-админу (сотруднику) → 403", st == 403, s)
    st, s, _ = call("GET", "/claims/summary?years=3", who="админ")
    g = next((x for x in s.get("items", []) if x["product_code"] == code and x["class_code"] == "8"), None)
    ok("сводка: продукт × класс", st == 200 and g is not None, s)
    ok("сводка: случаи, заявлено, выплачено, средняя", g["cases"] == 2 and g["claimed"] == 25_000_000
       and g["paid"] == 18_000_000 and g["avg_paid"] == 18_000_000, g)
    ok("сводка: убыточность = выплачено / премия (премия договора один раз)",
       g["premium"] == 4_000_000 and abs(g["loss_ratio"] - 18_000_000 / 4_000_000) < 1e-9, g)
    st, lst, _ = call("GET", "/claims?limit=5", who="админ")
    row = next(x for x in lst if x["external_no"] == "ТЕСТ-У-1")
    ok("GET /claims отдаёт product_code и contract_no строки", row["product_code"] == code
       and row["contract_no"] == "ТД-1", row)

    # калибровка видит новые строки по class_code
    with db.tx() as con:
        e = calibration.exposure_and_losses(con, (t - timedelta(days=365 * 3)).isoformat(), t.isoformat())
    g8 = e["groups"].get(("8", "склад")) or {}
    gna = e["groups"].get(("8", calibration.NO_OBJECT_TYPE)) or {}
    ok("калибровка учитывает загруженные случаи (класс 8, склад)", g8.get("m", 0) >= 1
       and g8.get("losses", 0) >= 18_000_000, g8)
    ok("случай без вида объекта — группа «не указан»", gna.get("m", 0) >= 1, gna)


def check_yearly(code: str):
    """Итоги года по продукту — ручной ввод в админке (06.10.2026): GET/PUT /claims/yearly."""
    y = date.today().year - 1
    body = {"year": y, "rows": [{"product_code": code, "cases": 3, "amount": 1000}]}
    st, r, _ = call("PUT", "/claims/yearly", body, who="сотрудник")
    ok("PUT /claims/yearly сотруднику закрыт", st in (401, 403), (st, r))
    st, r, _ = call("PUT", "/claims/yearly", {"year": y, "rows": [{"product_code": "НЕТ-ТАКОГО", "cases": 1, "amount": 5}]},
                    who="админ")
    ok("неизвестный продукт → 422 с понятной строкой", st == 422 and "нет в справочнике" in json.dumps(r, ensure_ascii=False), (st, r))
    st, r, _ = call("PUT", "/claims/yearly", body, who="админ")
    ok("PUT /claims/yearly: 3 случая записаны", st == 200 and r["cases"] == 3 and r["event_date"] == f"{y}-12-31", (st, r))
    st, d, _ = call("GET", f"/claims/yearly?years={y}", who="админ")
    m = [x for x in d["years"][0]["manual"] if x["product_code"] == code]
    ok("GET /claims/yearly: число случаев и сумма как введены (остаток округления сходится)",
       st == 200 and m and m[0]["cases"] == 3 and abs(m[0]["amount"] - 1000) < 1e-9, d)
    st, r, _ = call("PUT", "/claims/yearly", {"year": y, "rows": [{"product_code": code, "cases": 2, "amount": 500}]}, who="админ")
    st, d, _ = call("GET", f"/claims/yearly?years={y}", who="админ")
    m = [x for x in d["years"][0]["manual"] if x["product_code"] == code]
    ok("повторное сохранение года заменяет ручные итоги, а не добавляет", r["replaced"] == 3 and m[0]["cases"] == 2
       and abs(m[0]["amount"] - 500) < 1e-9, (r, m))
    st, s, _ = call("GET", "/claims/summary?years=3", who="админ")
    ok("сводка видит ручные итоги", st == 200 and any(x["product_code"] == code and x["cases"] >= 2 for x in s["items"]), s["items"][:3])
    st, r, _ = call("PUT", "/claims/yearly", {"year": date.today().year + 1, "rows": []}, who="админ")
    ok("будущий год → 422", st == 422, (st, r))


def main():
    with temp_db("surveyor-test-admin-manage.db"):
        db.ensure_schema()
        db.ensure_schema()               # миграция идемпотентна
        setup()
        check_staff()
        check_templates()
        code = check_products()
        check_claims(code)
        check_yearly(code)
    print(f"\nПройдено проверок: {len(PASSED)}")


if __name__ == "__main__":
    main()
