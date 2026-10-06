"""
Задача 166 (22.09.2026): шаблон отчётности портфеля в Excel, дашборд НАПП на /stats, выход из админки.

Запуск из корня проекта (живой сервер не нужен, рабочая база не меняется — временная копия):
    set PYTHONIOENCODING=utf-8
    sandbox\\.venv\\Scripts\\python.exe tests\\test_portfolio_template.py

Что проверяется:
  1. шаблон: четыре листа, обязательные колонки выделены, выпадающие списки, справочники из базы;
  2. заполненный шаблон (2 договора, 1 убыток) разбирается без ошибок, сохраняется и попадает в сводку;
     ИНН физлица не хранится;
  3. кривой файл: ошибки по строкам и колонкам, загрузка не делается, в базе ничего не появилось;
  4. через сервер: GET /portfolio/template.xlsx, POST /portfolio/preview и /portfolio/import, сводка загрузки;
  5. GET /stats — 200 и блоки дашборда; /tg/me?view=user у админа — меню и права как у сотрудника;
  6. кнопка «Выйти из админки» на /admin/deploy; на /admin и /admin/hub — «Выход» (06.10.2026).
"""
import asyncio
import io
import json as _json
import os
import sys
import tempfile
from datetime import date
from pathlib import Path
from urllib.parse import urlencode

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

os.environ.pop("SURVEYOR_DEV", None)          # guard проверяем по-настоящему

import openpyxl  # noqa: E402

from tmpdb import temp_db  # noqa: E402
from app import auth, db, tgbot  # noqa: E402
from app import portfolio as pf  # noqa: E402
from app.main import app  # noqa: E402

passed, failed = 0, 0
TOKENS = {}
PEOPLE = [("тест-портфель-сотрудник", "сотрудник"), ("тест-портфель-админ", "админ")]
EMP, ADM = PEOPLE[0][0], PEOPLE[1][0]


def ok(name, cond, extra=""):
    global passed, failed
    if cond:
        passed += 1
        print("  ок  ", name)
    else:
        failed += 1
        print("  ПЛОХО", name, str(extra)[:600])


def call(method, path, body=b"", params=None, who=None, ctype="application/json", headers=None):
    query = urlencode(params or {}, encoding="utf-8")
    if isinstance(body, (dict, list)):
        body = _json.dumps(body, ensure_ascii=False).encode()
    hdrs = [(b"host", b"test"), (b"content-type", ctype.encode()), (b"content-length", str(len(body)).encode())]
    if who:
        hdrs.append((b"cookie", f"sid={TOKENS[who]}".encode()))
    for k, v in (headers or {}).items():
        hdrs.append((k.lower().encode(), v.encode()))
    scope = {"type": "http", "asgi": {"version": "3.0"}, "http_version": "1.1", "method": method,
             "scheme": "http", "path": path, "raw_path": path.encode(), "root_path": "",
             "query_string": query.encode(), "headers": hdrs,
             "client": ("127.0.0.1", 0), "server": ("test", 80)}
    out = {"status": None, "chunks": [], "headers": {}}

    async def receive():
        return {"type": "http.request", "body": body, "more_body": False}

    async def send(msg):
        if msg["type"] == "http.response.start":
            out["status"] = msg["status"]
            out["headers"] = {k.decode().lower(): v.decode() for k, v in msg.get("headers", [])}
        elif msg["type"] == "http.response.body":
            out["chunks"].append(msg.get("body") or b"")

    asyncio.run(app(scope, receive, send))
    raw = b"".join(out["chunks"])
    if "json" in out["headers"].get("content-type", ""):
        return out["status"], _json.loads(raw.decode("utf-8"))
    return out["status"], raw


def multipart(name: str, data: bytes):
    b = "----surveyortest166"
    body = (f"--{b}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"{name}\"\r\n"
            "Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n").encode("utf-8")
    body += data + f"\r\n--{b}--\r\n".encode()
    return body, f"multipart/form-data; boundary={b}"


def setup():
    db.ensure_schema()
    with db.tx() as con:
        ts = db.now()
        for login, role in PEOPLE:
            cur = con.execute("INSERT INTO users (login, full_name, role, branch, password_hash, salt,"
                              " status, created_at) VALUES (?,?,?,?,?,?,?,?)",
                              (login, "Тест " + role, role, "тест-портфель", "x", "y", "активен", ts))
            u = db.rows(con, "SELECT * FROM users WHERE id=?", cur.lastrowid)[0]
            TOKENS[login], _ = auth.create_session(con, u, "127.0.0.1", "test")


# ---------- заполнение шаблона ----------

def headers_of(ws):
    return [c.value for c in ws[1]]


def put(ws, row: int, values: dict):
    """values: поле → значение; колонка ищется по заголовку шаблона."""
    heads = [str(h or "").rstrip(" *") for h in headers_of(ws)]
    for f, v in values.items():
        title = next(c[1] for c in (pf.PORTFOLIO_COLUMNS if ws.title == "Портфель" else pf.CLAIM_COLUMNS) if c[0] == f)
        ws.cell(row=row, column=heads.index(title) + 1, value=v)


GOOD = [
    {"external_no": "ТЕСТ/0807/26/0001", "date_signed": date(2026, 8, 25), "date_from": date(2026, 9, 1),
     "date_to": date(2027, 8, 31), "product": "0807", "class_code": "8", "branch": "Ташкентский",
     "region": "город Ташкент", "holder_type": "юрлицо", "holder_inn": "301234567", "object_type": "Склад",
     "sum_insured": 4_200_000_000, "value_amount": 4_200_000_000, "premium": 12_600_000, "premium_paid": 12_600_000,
     "franchise": 42_000_000, "rate": 0.3, "seismic_zone": 8, "currency": "UZS", "status": "действует",
     "claims_count": 1, "claims_claimed": 85_000_000, "claims_paid": 80_000_000, "last_loss_date": date(2026, 12, 3)},
    {"external_no": "ТЕСТ/0311/26/0002", "date_from": date(2026, 9, 10), "date_to": date(2027, 9, 9),
     "product": "0311", "branch": "Самаркандский", "region": "Самаркандская область", "holder_type": "физлицо",
     "object_type": "Легковой", "sum_insured": 146_100_000, "premium": 3_900_000, "premium_paid": 2_000_000,
     "rate": 2.6694, "status": "действует", "claims_count": 0},
]
CLAIM = {"contract_no": "ТЕСТ/0807/26/0001", "claim_no": "У-1", "event_date": date(2026, 12, 3),
         "reported_date": date(2026, 12, 5), "claimed": 85_000_000, "paid": 80_000_000, "status": "оплачен",
         "cause": "Пожар"}


def filled_template(folder: Path, bad: bool = False) -> Path:
    wb = openpyxl.load_workbook(io.BytesIO(pf.template_bytes()))
    ws, wc = wb["Портфель"], wb["Убытки"]
    for i, row in enumerate(GOOD, start=2):
        put(ws, i, row)
    put(wc, 2, CLAIM)
    if bad:
        # строка 4: дата текстом-ошибкой, сумма словом, ИНН у физлица, неизвестный продукт, чужой статус
        put(ws, 4, {"external_no": "ТЕСТ/9999/26/0003", "date_from": "31.02.2026", "date_to": date(2027, 1, 1),
                    "product": "9999", "region": "Марс", "holder_type": "физлицо", "holder_inn": "123456789",
                    "sum_insured": "много", "premium": 100, "status": "активен", "seismic_zone": 12})
        # строка 5: повтор номера и окончание раньше начала
        put(ws, 5, {"external_no": "ТЕСТ/0807/26/0001", "date_from": date(2026, 5, 1), "date_to": date(2026, 4, 1),
                    "product": "0807", "region": "Бухара", "holder_type": "юрлицо", "sum_insured": 10,
                    "premium": 1, "status": "действует"})
        put(wc, 3, {"contract_no": "НЕТ-ТАКОГО", "event_date": date(2026, 10, 1), "claimed": 5, "status": "оплачен"})
    p = folder / ("кривой.xlsx" if bad else "портфель_тест.xlsx")
    wb.save(p)
    return p


# ---------- 1. шаблон ----------

def check_template():
    print("1. Шаблон")
    wb = openpyxl.load_workbook(io.BytesIO(pf.template_bytes()))
    ok("четыре листа по порядку", wb.sheetnames == ["Портфель", "Убытки", "Инструкция", "Справочники"], wb.sheetnames)
    ws = wb["Портфель"]
    heads = headers_of(ws)
    ok("колонок на «Портфеле» — как в PORTFOLIO_COLUMNS", len(heads) == len(pf.PORTFOLIO_COLUMNS), heads)
    req = [c[1] + " *" for c in pf.PORTFOLIO_COLUMNS if c[3]]
    ok("обязательные помечены звёздочкой", all(h in heads for h in req), heads)
    fills = {ws.cell(row=1, column=heads.index(h) + 1).fill.fgColor.rgb[-6:] for h in req}
    ok("обязательные выделены цветом", fills == {pf.FILL_REQ}, fills)
    ok("есть выпадающие списки", any(dv.type == "list" for dv in ws.data_validations.dataValidation),
       [dv.type for dv in ws.data_validations.dataValidation])
    ok("у каждой колонки — подсказка в примечании", all(ws.cell(row=1, column=j).comment for j in range(1, len(heads) + 1)))
    ok("на «Убытках» — 8 колонок", len(headers_of(wb["Убытки"])) == len(pf.CLAIM_COLUMNS), headers_of(wb["Убытки"]))
    wr = wb["Справочники"]
    with db.tx() as con:
        n_prod = con.execute("SELECT COUNT(*) FROM products").fetchone()[0]
    prods = [c.value for c in wr["A"][1:] if c.value]
    ok("продукты в справочнике — из базы", len(prods) == n_prod and "0807" in prods, (len(prods), n_prod))
    regions = [c.value for c in wr["F"][1:] if c.value]
    ok("регионы — 14 из базы", len(regions) == 14 and "город Ташкент" in regions, regions)
    ok("PD: в шаблоне нет колонки ФИО", not any("фио" in str(h).lower() for h in heads), heads)
    wbu = openpyxl.load_workbook(io.BytesIO(pf.template_bytes("uz")))
    ok("шаблон по-узбекски: лист Portfel и узбекские заголовки",
       "Portfel" in wbu.sheetnames and "Shartnoma raqami *" in headers_of(wbu["Portfel"]), wbu.sheetnames)
    m = pf.map_columns(headers_of(wbu["Portfel"]))
    ok("узбекские заголовки узнаются все", len(m["index"]) == len(pf.PORTFOLIO_COLUMNS),
       sorted(set(c[0] for c in pf.PORTFOLIO_COLUMNS) - set(m["index"])))


# ---------- 2–3. разбор без сервера ----------

def check_import(tmp: Path):
    print("2. Заполненный шаблон")
    path = filled_template(tmp)
    a = pf.analyze_file(path)
    ok("режим «шаблон»", a["template"] and a["sheet"] == "Портфель", a["mode"])
    ok("ошибок нет", not a["errors"], a["errors"])
    ok("найдены все колонки шаблона", len(a["mapping"]) == len(pf.PORTFOLIO_COLUMNS), a["mapping"])
    ok("2 договора и 1 убыток", len(a["contracts"]) == 2 and len(a["claims"]) == 1, (len(a["contracts"]), len(a["claims"])))
    out = pf.import_file(path, path.name)
    ok("загружено", out["batch_id"] > 0 and out["summary"]["rows_total"] == 2, out.get("summary"))
    ov = out["overview"]
    t = ov["totals"]
    ok("сводка: 2 договора, сумма и премия", t["contracts"] == 2 and t["sum_insured"] == 4_200_000_000 + 146_100_000
       and t["premium"] == 12_600_000 + 3_900_000, t)
    ok("сводка: убытки с листа «Убытки»", t["claims"] == 1 and t["claims_paid"] == 80_000_000
       and ov["claims_source"] == "лист «Убытки»", (t, ov["claims_source"]))
    ok("по классам: 8 и 3", {b["key"] for b in ov["by_class"]} == {"8", "3"}, ov["by_class"])
    ok("по регионам: Ташкент и Самарканд",
       {b["key"] for b in ov["by_region"]} == {"город Ташкент", "Самаркандская область"}, ov["by_region"])
    with db.tx() as con:
        rows = {r["external_no"]: r for r in db.rows(con, "SELECT * FROM portfolio_contracts WHERE batch_id=?", out["batch_id"])}
        cl = db.rows(con, "SELECT * FROM portfolio_claims WHERE batch_id=?", out["batch_id"])
        again = pf.batch_overview(con, out["batch_id"])
    ok("ИНН юрлица сохранён, у физлица — пусто",
       rows["ТЕСТ/0807/26/0001"]["holder_inn"] == "301234567" and rows["ТЕСТ/0311/26/0002"]["holder_inn"] is None, rows)
    ok("сейсмозона, статус, валюта — в базе", rows["ТЕСТ/0807/26/0001"]["seismic_zone"] == 8
       and rows["ТЕСТ/0807/26/0001"]["status"] == "действует" and rows["ТЕСТ/0311/26/0002"]["currency"] == "UZS", rows)
    ok("убыток в portfolio_claims", len(cl) == 1 and cl[0]["status"] == "оплачен" and cl[0]["paid"] == 80_000_000, cl)
    ok("сводка из базы совпадает", again["totals"] == ov["totals"], (again["totals"], ov["totals"]))
    r0 = next(r for r in out["rows"] if r["external_no"] == "ТЕСТ/0807/26/0001")
    ok("движок проверил договор (класс 8, объект «Склад»)", r0["verdict"] != "не проверен" and r0.get("object_type") == "Склад", r0)

    print("3. Кривой файл")
    bad = filled_template(tmp, bad=True)
    a = pf.analyze_file(bad)
    errs = a["errors"]
    where = {(e["sheet"], e["row"], e["column"]) for e in errs}
    has = lambda row, col_part, msg_part: any(e["row"] == row and col_part in e["column"] and msg_part in e["message"] for e in errs)
    ok("дата 31.02 — ошибка в строке 4", has(4, "Дата начала", "дата не распознана"), errs)
    ok("«много» — не число", has(4, "Страховая сумма", "не число"), errs)
    ok("ИНН у физлица — ошибка, значение не выводится",
       has(4, "ИНН", "персональные данные") and not any("123456789" in e["message"] for e in errs), errs)
    ok("продукт 9999 — не из справочника", has(4, "Код продукта", "не найден"), errs)
    ok("регион «Марс» — не из списка", has(4, "Регион", "не из списка"), errs)
    ok("статус «активен» — ошибка", has(4, "Статус", "статус"), errs)
    ok("сейсмозона 12 — ошибка", has(4, "Сейсмо", "от 6 до 10"), errs)
    ok("повтор номера договора", has(5, "Номер договора", "повторяется"), errs)
    ok("окончание раньше начала", has(5, "Дата окончания", "не позже"), errs)
    ok("«Бухара» узнана как область", not has(5, "Регион", ""), [e for e in errs if e["row"] == 5])
    ok("убыток по несуществующему договору", any(e["sheet"] == "Убытки" and e["row"] == 3 and "нет на листе" in e["message"] for e in errs), errs)
    ok("у каждой ошибки есть лист, строка, колонка", all(all(x) for x in where), where)
    with db.tx() as con:
        n0 = con.execute("SELECT COUNT(*) FROM portfolio_batches").fetchone()[0]
    try:
        pf.import_file(bad, bad.name)
        ok("кривой шаблон не загружается", False)
    except pf.RowErrors as e:
        ok("кривой шаблон не загружается: RowErrors со списком", len(e.errors) == len(errs), str(e))
    with db.tx() as con:
        n1 = con.execute("SELECT COUNT(*) FROM portfolio_batches").fetchone()[0]
    ok("в базе новой загрузки нет", n1 == n0, (n0, n1))


# ---------- 4–6. через сервер ----------

def check_api(tmp: Path):
    print("4. Через сервер")
    st, raw = call("GET", "/portfolio/template.xlsx", who=ADM)
    ok("GET /portfolio/template.xlsx — 200 и xlsx", st == 200 and raw[:2] == b"PK", st)
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    ok("скачанный файл открывается openpyxl", "Портфель" in wb.sheetnames, wb.sheetnames)
    st, _ = call("GET", "/portfolio/template.xlsx")
    ok("без входа шаблон не отдаётся", st in (401, 303, 302, 307), st)

    body, ct = multipart("портфель_api.xlsx", filled_template(tmp).read_bytes())
    st, j = call("POST", "/portfolio/preview", body, who=ADM, ctype=ct)
    ok("предпросмотр — 200, ошибок нет, можно сохранить", st == 200 and j["errors_total"] == 0 and j["can_import"], j)
    ok("предпросмотр: имён страхователей нет", all("policyholder" not in r for r in j.get("rows", [])))
    with db.tx() as con:
        n0 = con.execute("SELECT COUNT(*) FROM portfolio_batches").fetchone()[0]
    st, j = call("POST", "/portfolio/import", body, who=ADM, ctype=ct)
    ok("импорт — 200", st == 200 and j.get("batch_id"), (st, j if st != 200 else ""))
    bid = j.get("batch_id")
    with db.tx() as con:
        n1 = con.execute("SELECT COUNT(*) FROM portfolio_batches").fetchone()[0]
    ok("предпросмотр ничего не сохранил, импорт — одну загрузку", n1 == n0 + 1, (n0, n1))
    st, o = call("GET", f"/portfolio/batches/{bid}/overview", who=ADM)
    ok("сводка загрузки: 2 договора, 1 убыток", st == 200 and o["overview"]["totals"]["contracts"] == 2
       and o["overview"]["totals"]["claims"] == 1, o)
    st, o = call("GET", "/portfolio/overview", who=ADM)
    ok("/portfolio/overview — последняя загрузка", st == 200 and o["batch"]["id"] == bid, o.get("batch"))
    st, lst = call("GET", "/portfolio/batches", who=ADM)
    ok("в списке загрузок есть число убытков", st == 200 and lst[0]["claims_rows"] == 1, lst[:1])

    body, ct = multipart("кривой.xlsx", filled_template(tmp, bad=True).read_bytes())
    st, j = call("POST", "/portfolio/preview", body, who=ADM, ctype=ct)
    ok("кривой: предпросмотр 200, ошибки по строкам, сохранить нельзя",
       st == 200 and j["errors_total"] >= 10 and not j["can_import"] and j["errors"][0]["row"], j.get("errors_total"))
    st, j = call("POST", "/portfolio/import", body, who=ADM, ctype=ct)
    ok("кривой: импорт — 400 со списком ошибок", st == 400 and isinstance(j["detail"], dict)
       and len(j["detail"]["errors"]) >= 10, (st, str(j)[:300]))

    print("5. Дашборд /stats и /tg/me?view=user")
    st, html = call("GET", "/stats", who=ADM)
    html = html.decode("utf-8") if isinstance(html, bytes) else str(html)
    blocks = ["id=\"tabDash\"", "id=\"kpis\"", "id=\"chDyn\"", "id=\"chLoss\"", "id=\"chMix\"", "id=\"chTop\"",
              "id=\"chReg\"", "id=\"fDate\"", "id=\"fCls\"", "id=\"fReg\"", "id=\"tabBuilder\"", ">Конструктор<",
              "napp.uz", "/theme.js", "/i18n.js"]
    ok("GET /stats — 200", st == 200, st)
    ok("все блоки дашборда на месте", all(b in html for b in blocks), [b for b in blocks if b not in html])
    ok("внешних библиотек графиков нет", "cdn.jsdelivr" not in html and "chart.js" not in html.lower())
    st, html = call("GET", "/stats", params={"embed": 1}, who=ADM)
    ok("/stats?embed=1 для админки — 200", st == 200, st)

    st, me = call("GET", "/tg/me", params={"view": "user"}, who=ADM)
    emp_nav = [k for k, _ in tgbot.NAV_BASE]
    ok("view=user: меню как у сотрудника", [n["key"] for n in me["nav"]] == emp_nav, me.get("nav"))
    ok("view=user: can_edit пуст, is_admin=false, role=сотрудник",
       me["can_edit"] == [] and me["user"]["is_admin"] is False and me["user"]["role"] == "сотрудник", me["user"])
    ok("view=user: view, real_role, admin_available", me["view"] == "user" and me["real_role"] == "админ"
       and me["admin_available"] is True, me)
    st, me_e = call("GET", "/tg/me", who=EMP)
    ok("права в view=user = права сотрудника", me["rights"] == me_e["rights"], (me["rights"], me_e["rights"]))
    st, me2 = call("GET", "/tg/me", who=ADM, headers={"X-View": "user"})
    ok("заголовок X-View: user — то же", [n["key"] for n in me2["nav"]] == emp_nav and me2["view"] == "user", me2.get("nav"))
    st, me3 = call("GET", "/tg/me", who=ADM)
    ok("без view — полное админское меню", "settings" in [n["key"] for n in me3["nav"]] and me3["view"] == "full"
       and me3["can_edit"], me3.get("nav"))
    st, _ = call("GET", "/audit", params={"limit": 1}, who=ADM, headers={"X-View": "user"})
    ok("права админа на сервере не изменились (/audit с X-View: user — 200)", st == 200, st)
    st, me4 = call("GET", "/tg/me", params={"view": "user"}, who=EMP)
    ok("сотрудник с view=user: admin_available=false", me4["admin_available"] is False and me4["view"] == "user", me4)

    print("6. Кнопка «Выйти из админки»")
    # 06.10.2026 (записка заказчика): админка /admin — отдельный сайт; режима mode=user в её шапке нет,
    # вместо него «Выход» из учётной записи. Кнопка в мини-апп осталась на /admin/deploy.
    for path in ("/admin/hub", "/admin"):
        st, html = call("GET", path, who=ADM)
        html = html.decode("utf-8") if isinstance(html, bytes) else str(html)
        ok(f"{path}: без mode=user, есть «Выход» через /auth/logout",
           st == 200 and "/tg?mode=user" not in html and "/auth/logout" in html and ">Выход<" in html, st)
    st, html = call("GET", "/admin/deploy", who=ADM)
    html = html.decode("utf-8") if isinstance(html, bytes) else str(html)
    ok("/admin/deploy: кнопка ведёт на /tg?mode=user", st == 200 and "/tg?mode=user" in html and "Выйти из админки" in html, st)
    st, html = call("GET", "/admin/deploy", params={"embed": 1}, who=ADM)
    html = html.decode("utf-8") if isinstance(html, bytes) else str(html)
    ok("во встроенном окне (embed=1) второй кнопки нет", "exitAdmin" not in html, "")


def main():
    with temp_db("surveyor-portfolio-template.db"), tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        pf.UPLOAD_DIR = tmp / "uploads"          # загрузки — во временную папку, не в data/
        setup()
        check_template()
        check_import(tmp)
        check_api(tmp)
    print(f"\nИтого: пройдено {passed}, не пройдено {failed}")
    if failed:
        sys.exit(1)


if __name__ == "__main__":
    main()
