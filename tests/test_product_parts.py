"""
Части продуктов (ТЗ разработчику 13.10.2026, день 2.1–2.3): миграция из тарифной политики, ставка части в акте,
тип клиента, импорт Excel и API. Всё — на ВРЕМЕННОЙ копии базы (tests/tmpdb.py).

Запуск из корня проекта:
    set PYTHONIOENCODING=utf-8
    sandbox\\.venv\\Scripts\\python.exe tests\\test_product_parts.py
"""
import asyncio
import io
import json
import os
import sys
from datetime import date, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))
os.environ.pop("SURVEYOR_DEV", None)

import openpyxl                                        # noqa: E402

from tmpdb import temp_db                              # noqa: E402
from app import act_engine as ae, auth, db, min_rates, product_parts as pp   # noqa: E402
from app.main import app                               # noqa: E402

PASSED = []
TOK = {}
TODAY = date.today()


def ok(name, cond, hint=""):
    assert cond, f"ПРОВАЛ: {name} {hint}"
    PASSED.append(name)
    print("  ✓", name)


def call(method, path, body=None, who=None, raw=None, ctype="application/json"):
    payload = raw if raw is not None else (json.dumps(body, ensure_ascii=False).encode() if body is not None else b"")
    path, _, query = path.partition("?")
    hdrs = [(b"host", b"test"), (b"content-type", ctype.encode()), (b"content-length", str(len(payload)).encode())]
    if who:
        hdrs.append((b"authorization", f"Bearer {TOK[who]}".encode()))
    scope = {"type": "http", "asgi": {"version": "3.0"}, "http_version": "1.1", "method": method,
             "scheme": "http", "path": path, "raw_path": path.encode(), "root_path": "",
             "query_string": query.encode(), "headers": hdrs, "client": ("203.0.113.9", 0), "server": ("test", 80)}
    out = {"status": None, "chunks": []}

    async def receive():
        return {"type": "http.request", "body": payload, "more_body": False}

    async def send(msg):
        if msg["type"] == "http.response.start":
            out["status"] = msg["status"]
        elif msg["type"] == "http.response.body":
            out["chunks"].append(msg.get("body") or b"")

    asyncio.run(app(scope, receive, send))
    data = b"".join(out["chunks"])
    try:
        return out["status"], json.loads(data.decode("utf-8"))
    except (ValueError, UnicodeDecodeError):
        return out["status"], data


def upload(path, data, who):
    b = "----parts-boundary"
    body = (f"--{b}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"части.xlsx\"\r\n"
            "Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n").encode() \
        + data + f"\r\n--{b}--\r\n".encode()
    return call("POST", path, raw=body, who=who, ctype=f"multipart/form-data; boundary={b}")


def xlsx(rows) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def setup():
    with db.tx() as con:
        for login, role in (("тест-части-админ", "админ"), ("тест-части-сотрудник", "сотрудник")):
            cur = con.execute(
                "INSERT INTO users (login, full_name, phone, role, branch, password_hash, salt, status, created_at)"
                " VALUES (?,?,?,?,?,?,?,?,?)",
                (login, "Тест " + role, "", role, "тест", "x" * 64, "0" * 32, auth.STATUS_ACTIVE, db.now()))
            u = db.rows(con, "SELECT * FROM users WHERE id=?", cur.lastrowid)[0]
            TOK[role], _ = auth.create_session(con, u, ip="203.0.113.9", user_agent="test")


def act_data(aid: str) -> dict:
    with db.tx() as con:
        r = db.rows(con, "SELECT act_json FROM acts WHERE id=?", aid)
    return json.loads(r[0]["act_json"])["data"]


# --------------------------------------------------------------------------- #

def check_migration():
    print("Миграция из тарифной политики")
    with db.tx() as con:
        con.execute("DELETE FROM product_parts")           # копия базы: проверяем миграцию с нуля
        con.execute("DELETE FROM product_part_versions")
        res = pp.migrate_from_policy(con, "тест")
        again = pp.migrate_from_policy(con, "тест")
        n_prod = con.execute("SELECT COUNT(*) FROM products").fetchone()[0]
        n_pc = con.execute("SELECT COUNT(*) FROM product_classes").fetchone()[0]
        n_payer = con.execute("SELECT COUNT(*) FROM min_rates m JOIN tariff_versions v ON v.id=m.tariff_version_id "
                              "WHERE v.name=? AND m.payer_type IS NOT NULL", (pp.POLICY_NAME,)).fetchone()[0]
        n_rows = con.execute("SELECT COUNT(*) FROM product_parts").fetchone()[0]
        p833 = pp.parts_on(con, "0833", TODAY)
        p309 = pp.parts_on(con, "0309", TODAY)
        fiz = pp.rate_part(con, "0309", "3", "fiz", TODAY)
        yur = pp.rate_part(con, "0309", "3", "yur", TODAY)
        none_ = pp.rate_part(con, "0309", "3", None, TODAY)
        neg = pp.rate_part(con, "0824", "8", None, TODAY)
        stat = db.rows(con, "SELECT code FROM products WHERE pricing_mode='нормативный акт' LIMIT 1")[0]["code"]
        stat_cls = db.rows(con, "SELECT class_code FROM product_classes WHERE product_code=?", stat)[0]["class_code"]
        stat_part = pp.rate_part(con, stat, stat_cls, None, TODAY)
        cur_v = min_rates.current_version(con, "компания")
    ok(f"все продукты разобраны: {res['products']} из {n_prod}", res["products"] == n_prod, res)
    ok(f"частей = строк product_classes ({n_pc})", res["parts"] == n_pc, res["parts"])
    ok(f"строк = частей + ставки по типу клиента ({n_pc} + {n_payer // 2})",
       res["rows"] == n_rows == n_pc + n_payer // 2, (res["rows"], n_rows, n_payer))
    ok("повторный запуск ничего не добавляет", again["rows"] == 0 and again["skipped"] == n_prod, again)
    ok("спорные разборы перечислены (0833 — общая ставка, 0327 — ставки по типу ТС)",
       {"0833", "0327"} <= {d["product_code"] for d in res["disputed"]}, [d["product_code"] for d in res["disputed"]])
    parts = p833["parts"]
    ok("0833 — две части: кл. 8 и кл. 13", [(p["part_no"], p["class_code"]) for p in parts] == [(1, "8"), (2, "13")],
       parts)
    ok("у каждой части 0833 своя базовая ставка и минимум",
       all(p["base_rate_pct"] is not None and p["min_rate_pct"] is not None for p in parts), parts)
    ok("0833: база суммы — стоимость у имущества, лимит у ответственности",
       [p["sum_basis"] for p in parts] == ["value", "limit"], parts)
    ok("версия частей — тарифная политика (valid_from, source, created_at)",
       p833["version"]["source"] == "policy" and p833["version"]["valid_from"] == "2025-09-23"
       and p833["version"]["created_at"], p833["version"])
    ok("0309: две строки части 1 — физлицо и юрлицо",
       sorted((p["client_type"], p["min_rate_pct"]) for p in p309["parts"]) == [("fiz", 1.2), ("yur", 1.5)],
       p309["parts"])
    ok("0309 для физлица — 1,2 %, для юрлица — 1,5 %",
       fiz and yur and fiz["base_rate_pct"] == 1.2 and yur["base_rate_pct"] == 1.5, (fiz, yur))
    ok("тип клиента не задан — строки для всех нет, часть не выбирается (прежний путь)", none_ is None, none_)
    ok("«по согласованию» и обязательный вид: у части нет ставки → прежний путь",
       neg is None and stat_part is None, (neg, stat_part))
    ok("версия частей не подменяет версию тарифа уровня (current_version)", cur_v == 1, cur_v)


def check_engine_same():
    """Строки политики дают ровно тот же расчёт, что прежний путь: все продукты, классы, типы клиента, уровни."""
    print("Прежние ставки и премии не изменились")
    S, n, diffs = 1_000_000_000, 0, []
    with db.tx() as con:
        ref = db.load_reference(con)
        prods = db.rows(con, "SELECT code, name, pricing_mode, rate_text FROM products ORDER BY code")
        for p in prods:
            classes = list(ref.product_classes.get(p["code"]) or [])
            for ct in (None, "fiz", "yur"):
                payer = pp.PAYER_OF.get(ct)
                for i, cls in enumerate(classes):
                    part = pp.rate_part(con, p["code"], cls, ct, TODAY)
                    for lvl in ae.LEVELS:
                        if i == 0:
                            a = ae.rate(ref, p, cls, lvl, S, 200, None, payer)
                            b = ae.rate(ref, p, cls, lvl, S, 200, None, payer, part=part)
                        else:
                            a = ae.part_rate(ref, p, cls, lvl, S, 200, None, payer, None, ae.class_min(ref, p, cls))
                            b = ae.part_rate(ref, p, cls, lvl, S, 200, None, payer, None,
                                             ae.class_min(ref, p, cls, part))
                        n += 1
                        ka = [a.get(k) for k in ("mode", "base_pct", "applied_pct", "min_pct", "premium", "rate_type")]
                        kb = [b.get(k) for k in ("mode", "base_pct", "applied_pct", "min_pct", "premium", "rate_type")]
                        if ka != kb or [h["code"] for h in a["how"]] != [h["code"] for h in b["how"]]:
                            diffs.append((p["code"], cls, ct, lvl, ka, kb))
    ok(f"{n} расчётов: ставка, минимум, премия и строки «как посчитано» те же", not diffs, diffs[:3])


def check_client_type():
    print("Тип клиента")
    st = ae.merge_settings({})
    ok("0833 «… физических лиц» → физлицо по продукту",
       pp.default_client_type({"code": "0833", "name": "Имущество и ГО физических лиц"}, st) == ("fiz", "product"))
    ok("0101 «НС юридических лиц» → юрлицо по продукту",
       pp.default_client_type({"code": "0101", "name": "НС юридических лиц"}, st) == ("yur", "product"))
    ok("0309 КАСКО — тип не определён → юрлицо по умолчанию",
       pp.default_client_type({"code": "0309", "name": "КАСКО «Standart»"}, st) == ("yur", "default"))
    ok("настройки client_type проходят проверку", not ae.check_settings({}), ae.check_settings({}))
    ok("ошибка в client_type ловится", ae.check_settings({"client_type": {"default": "x"}}))
    ok("оговорка работодателя у физлица убирается, у юрлица остаётся",
       ae.client_filter(["lb_limit", "lb_employer_payroll"], st, "fiz", "clauses") == ["lb_limit"]
       and ae.client_filter(["lb_limit", "lb_employer_payroll"], st, "yur", "clauses")
       == ["lb_limit", "lb_employer_payroll"])

    base = {"lang": "ru", "must": {"product_code": "0309", "sum_insured": 100_000_000, "object_value": 100_000_000,
                                   "region": "Ташкентская область"},
            "optional": {"losses_3y": {"count": 0, "small_count": 0, "amount": 0}, "term_days": 365}}
    out = {}
    for ct in ("fiz", "yur", None):
        b = json.loads(json.dumps(base))
        if ct:
            b["must"]["client_type"] = ct
        st_, a = call("POST", "/act/make", b, who="админ")
        ok(f"акт 0309 ({ct or 'по умолчанию'}) сформирован", st_ == 200 and a.get("ok"), (st_, str(a)[:300]))
        out[ct] = act_data(a["id"])
    f, y, d = out["fiz"], out["yur"], out[None]
    ok("0309 физлицо: база и минимум 1,2 %, премия = 100 млн × ставка",
       f["rate"]["base_pct"] == 1.2 and f["rate"]["min_pct"] == 1.2
       and f["rate"]["premium"] == round(f["rate"]["applied_pct"] / 100 * 100_000_000), f["rate"])
    ok("0309 юрлицо: база и минимум 1,5 %", y["rate"]["base_pct"] == 1.5 and y["rate"]["min_pct"] == 1.5, y["rate"])
    ok("строка тарифа — из части продукта (client_type в данных ставки)",
       f["rate"]["part"]["client_type"] == "fiz" and y["rate"]["part"]["client_type"] == "yur"
       and f["product_part"]["source"] == "policy", f["rate"].get("part"))
    ok("по умолчанию — юрлицо (1,5 %, как раньше: наибольшая из ставок)",
       d["must"]["client_type"] == "yur" and d["must"]["client_type_source"] == "default"
       and d["rate"]["min_pct"] == 1.5, d["must"])

    b = json.loads(json.dumps(base))
    b["must"]["client_type"] = "ip"
    st_, a = call("POST", "/act/make", b, who="админ")
    ok("неверный client_type → 422", st_ == 422, (st_, a))
    b["must"]["client_type"] = "fiz"
    b["optional"]["payer_type"] = "юр"
    st_, a = call("POST", "/act/make", b, who="админ")
    ok("client_type и payer_type расходятся → 422", st_ == 422, (st_, a))

    # ГО физлица (0833, части 8 и 13): ФОТ и оборот не принимаются, оговорок работодателя нет
    go = {"lang": "ru", "must": {"product_code": "0833", "sum_insured": 200_000_000, "object_value": 220_000_000,
                                 "region": "Ташкентская область"},
          "optional": {"losses_3y": {"count": 0, "small_count": 0, "amount": 0}}}
    b = json.loads(json.dumps(go))
    b["optional"]["class_fields"] = {"turnover_or_payroll": 500_000_000}
    st_, a = call("POST", "/act/make", b, who="админ")
    ok("0833 физлицо + ФОТ/оборот → 422 «только для юрлица»", st_ == 422 and "юрлиц" in json.dumps(a, ensure_ascii=False),
       (st_, a))
    st_, a = call("POST", "/act/make", go, who="админ")
    ok("акт 0833 (физлицо по продукту) сформирован", st_ == 200 and a.get("ok"), (st_, str(a)[:300]))
    D = act_data(a["id"])
    items = D["parts"]["items"]
    ok("0833: акт по двум частям, кл. 8 и 13", [p["class_code"] for p in items] == ["8", "13"],
       [p["class_code"] for p in items])
    ok("0833: у каждой части своя ставка и минимум из частей продукта",
       all(p["rate"]["min_pct"] == 0.1 and (p["class_min"].get("part") or {}).get("part_no") == p["index"]
           for p in items), [(p["rate"]["min_pct"], p["class_min"]) for p in items])
    ok("0833: премия договора = сумма премий частей",
       D["premium_final"]["amount"] == round(sum(p["premium"] for p in items)), D["premium_final"])
    codes = {c["code"] for c in D["clauses"]} | {c["code"] for p in items for c in p["clauses"]}
    ok("ГО физлица: нет оговорки работодателя", "lb_employer_payroll" not in codes, codes)
    meas = {m.get("code") for m in D["measures"]["items"]}
    ok("ГО физлица: нет мероприятия по охране труда работников", "gl_safety" not in meas, meas)
    ok("ГО физлица: ФОТ/оборот не числится среди недостающих полей",
       all(x["code"] != "turnover_or_payroll" for p in items for x in p.get("missing") or []))
    b = json.loads(json.dumps(go))
    b["must"]["client_type"] = "yur"
    b["optional"]["class_fields"] = {"turnover_or_payroll": 500_000_000}
    st_, a = call("POST", "/act/make", b, who="админ")
    ok("0833 юрлицо с ФОТ/оборотом — принято", st_ == 200 and a.get("ok"), (st_, str(a)[:300]))
    D = act_data(a["id"])
    codes = {c["code"] for c in D["clauses"]}
    ok("у юрлица оговорка работодателя остаётся", "lb_employer_payroll" in codes, codes)


def _row(code="0833", no=1, cls="8", ct="both", base=0.12, mn=0.1, share=0.7, mand="да", d=None):
    return [code, no, cls, "", ct, base, mn, "annual", "value" if cls in ("8", "9", "3") else "limit", share, "",
            "", mand, no, (d or TODAY).strftime("%d.%m.%Y"), "тест"]


def check_excel():
    print("Excel: шаблон, предпросмотр, отказ по каждой проверке, запись версией")
    st_, data = call("GET", "/reference/product-parts/template.xlsx")
    ok("шаблон — xlsx", st_ == 200 and isinstance(data, bytes) and data[:2] == b"PK")
    wb = openpyxl.load_workbook(io.BytesIO(data))
    head = [c.value for c in wb.worksheets[0][1]]
    ok("в шаблоне все поля части", head == pp.HEAD and "Инструкция" in wb.sheetnames, head)

    bad = {
        "ставка ниже минимума": ([pp.HEAD, _row(base=0.05, mn=0.1), _row(no=2, cls="13", share=0.3)],
                                 "ниже минимальной"),
        "сумма долей обязательных частей > 1": ([pp.HEAD, _row(share=0.8), _row(no=2, cls="13", share=0.4)],
                                                "больше 1"),
        "класс не из справочника": ([pp.HEAD, _row(cls="99"), _row(no=2, cls="13", share=0.3)],
                                    "нет в справочнике классов"),
        "дата в прошлом": ([pp.HEAD, _row(d=TODAY - timedelta(days=1)), _row(no=2, cls="13", share=0.3,
                                                                          d=TODAY - timedelta(days=1))],
                           "в прошлом"),
    }
    for name, (rows, word) in bad.items():
        st_, pv = upload("/reference/product-parts/import", xlsx(rows), "админ")
        ok(f"{name}: файл отклонён с понятным сообщением",
           st_ == 200 and pv["can_apply"] is False and pv["token"] is None and word in pv["detail"],
           (st_, pv.get("detail")))
    st_, out = call("POST", "/reference/product-parts/import/apply", {"token": ""}, who="админ")
    ok("применить без token (файл с ошибками его не получает) → 422", st_ == 422, (st_, out))
    st_, pv = upload("/reference/product-parts/import", xlsx([pp.HEAD, _row(), _row(no=2, cls="13", share=0.3)]),
                     "сотрудник")
    ok("импорт не администратором → 403", st_ == 403, st_)

    # корректный файл: доли 0,7 + 0,3 = 1; необязательная часть в сумму долей не входит
    good = [pp.HEAD, _row(base=0.12, mn=0.1, share=0.7), _row(no=2, cls="13", base=0.6, mn=0.5, share=0.3),
            _row(no=3, cls="9", base=0.2, mn=0.15, share=0.5, mand="нет")]
    st_, pv = upload("/reference/product-parts/import", xlsx(good), "админ")
    ok("корректный файл: предпросмотр без ошибок, есть token", st_ == 200 and pv["can_apply"] and pv["token"],
       (st_, pv.get("detail"), pv.get("messages")))
    ok("предпросмотр показывает прежний состав (было → станет)",
       pv["products"][0]["now"]["version"]["source"] == "policy" and len(pv["products"][0]["now"]["parts"]) == 2,
       pv["products"][0]["now"])
    ok("предупреждение: класс 9 не входит в состав продукта",
       any("не входит в состав" in w for w in pv["products"][0]["warnings"]), pv["products"][0]["warnings"])
    st_, out = call("POST", "/reference/product-parts/import/apply", {"token": pv["token"]}, who="админ")
    ok("применить → новая версия", st_ == 200 and out["ok"] and out["saved"][0]["parts"] == 3, (st_, out))
    vid = out["saved"][0]["tariff_version_id"]
    st_, cur = call("GET", "/reference/product-parts/0833")
    ok("GET на сегодня — новая версия (импорт), три части",
       st_ == 200 and cur["version"]["tariff_version_id"] == vid and cur["version"]["source"] == "import"
       and len(cur["parts"]) == 3, cur.get("version"))
    st_, old = call("GET", "/reference/product-parts/0833?date=2026-01-01")
    ok("GET на прошлую дату — версия тарифной политики (старые расчёты воспроизводятся)",
       st_ == 200 and old["version"]["source"] == "policy" and len(old["parts"]) == 2, old.get("version"))
    ok("история версий: политика и импорт", [v["source"] for v in cur["versions"]] == ["policy", "import"],
       cur["versions"])
    with db.tx() as con:
        part = pp.rate_part(con, "0833", "13", "fiz", TODAY)
        cv = min_rates.current_version(con, "компания")
    ok("ставка акта для кл. 13 — из новой версии (0,6 / 0,5)",
       part and part["base_rate_pct"] == 0.6 and part["min_rate_pct"] == 0.5, part)
    ok("версия частей не становится версией тарифа уровня", cv == 1, cv)
    st_, a = call("POST", "/act/make", {"lang": "ru", "must": {"product_code": "0833", "sum_insured": 200_000_000,
                                                               "object_value": 220_000_000,
                                                               "region": "Ташкентская область"},
                                        "optional": {"losses_3y": {"count": 0, "small_count": 0, "amount": 0}}},
                  who="админ")
    ok("акт 0833 после импорта сформирован", st_ == 200 and a.get("ok"), (st_, str(a)[:300]))
    D = act_data(a["id"])
    p13 = next(p for p in D["parts"]["items"] if p["class_code"] == "13")
    ok("часть кл. 13: база 0,6, минимум 0,5, источник — справочник частей",
       p13["rate"]["base_pct"] == 0.6 and p13["rate"]["min_pct"] == 0.5
       and p13["class_min"]["source"] == "product_parts", (p13["rate"]["base_pct"], p13["rate"]["min_pct"]))
    with db.tx() as con:
        ref = db.load_reference(con)
        prod = db.rows(con, "SELECT code, name, pricing_mode, rate_text FROM products WHERE code='0833'")[0]
        cm13 = ae.class_min(ref, prod, "13", pp.rate_part(con, "0833", "13", "fiz", TODAY))
        r13 = ae.part_rate(ref, prod, "13", "low", 1e8, 365, None, "физ", None, cm13)
        # прежняя формула (модель ставки выключена) — видна строка «как посчитано» о части
        r13o = ae.part_rate(ref, prod, "13", "low", 1e8, 365, None, "физ", {"rate_model": {"enabled": False}}, cm13)
    ok("движок: ставка части — из импорта (данные ставки хранят часть и её версию)",
       r13["part"]["source"] == "import" and r13["part"]["tariff_version_id"] == vid and r13["min_pct"] == 0.5,
       r13.get("part"))
    ok("строка «как посчитано» о части из справочника частей",
       r13o["how"][0]["code"] == "how_part_src" and r13o["applied_pct"] == 0.6, r13o["how"][:2])
    ok("строка «как посчитано» о части есть на трёх языках",
       all(ae is not None and __import__("app.act_texts", fromlist=["t"]).t("how_part_src", lg, code="0833", n=2,
                                                                           cls="13", base="0,6 %", min="0,5 %",
                                                                           date="2026-10-13") for lg in ("ru", "uz",
                                                                                                         "en")))

    # правка администратора в min_rates новее частей — она и действует (прежний путь)
    st_, out = call("PUT", "/reference/min-rates/0833", {"min_rate_pct": 0.2}, who="админ")
    ok("минимальная ставка 0833 изменена администратором", st_ == 200, (st_, out))
    with db.tx() as con:
        part8 = pp.rate_part(con, "0833", "8", "fiz", TODAY)
    ok("min_rates новее частей → часть не берётся, действует правка минимума", part8 is None, part8)


def check_put():
    print("API формы: PUT /reference/product-parts/{code}")
    body = {"valid_from": (TODAY + timedelta(days=30)).isoformat(),
            "parts": [{"part_no": 1, "class_code": "8", "client_type": "both", "base_rate_pct": 0.1,
                       "min_rate_pct": 0.2, "default_share": 0.6}]}
    st_, out = call("PUT", "/reference/product-parts/0806", body, who="админ")
    ok("ставка ниже минимума → 422 с текстом", st_ == 422 and "ниже минимальной" in out["detail"], (st_, out))
    st_, out = call("PUT", "/reference/product-parts/0806", body, who="сотрудник")
    ok("не администратор → 403", st_ == 403, st_)
    body["parts"] = [{"part_no": 1, "class_code": "8", "base_rate_pct": 0.15, "min_rate_pct": 0.1,
                      "default_share": 0.6, "perils": ["fire"]},
                     {"part_no": 2, "class_code": "13", "base_rate_pct": 0.5, "min_rate_pct": 0.5,
                      "default_share": 0.4, "sum_basis": "limit", "limit": 50_000_000, "deductible": 1}]
    st_, out = call("PUT", "/reference/product-parts/0806", body, who="админ")
    ok("корректная версия с будущей датой сохранена", st_ == 200 and out["ok"], (st_, out))
    st_, now_ = call("GET", "/reference/product-parts/0806")
    st2, fut = call("GET", f"/reference/product-parts/0806?date={body['valid_from']}")
    ok("до даты начала — прежние части, с даты — новые",
       now_["version"]["source"] == "policy" and fut["version"]["source"] == "admin"
       and fut["parts"][1]["limit"] == 50_000_000 and fut["parts"][0]["perils"] == ["fire"], (now_["version"],
                                                                                          fut["version"]))
    st_, out = call("GET", "/reference/product-parts/9999")
    ok("неизвестный продукт → 404", st_ == 404, st_)


def main():
    with temp_db():
        db.ensure_schema()
        setup()
        check_migration()
        check_engine_same()
        check_client_type()
        check_excel()
        check_put()
    print(f"\nхорошо {len(PASSED)}, плохо 0")


if __name__ == "__main__":
    main()
